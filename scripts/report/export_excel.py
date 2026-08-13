"""Export reproducible analysis tables to a multi-sheet Excel workbook."""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
PITCHERS = ROOT / "data" / "processed" / "pitcher_stats.csv"
TEAM = ROOT / "data" / "processed" / "team_financials.csv"
DID_RESULTS = ROOT / "data" / "processed" / "did_regression_results.csv"
PITCHER_RESULTS = ROOT / "data" / "processed" / "pitcher_regression_results.csv"
OUTPUT = ROOT / "reports" / "analysis_summary.xlsx"
BUCKET_LABELS = ["1) <92 mph", "2) 92-94 mph", "3) 94-96 mph", "4) 96+ mph"]


def main() -> None:
    pitchers = pd.read_csv(PITCHERS, encoding="utf-8-sig")
    team = pd.read_csv(TEAM, encoding="utf-8-sig")
    did = pd.read_csv(DID_RESULTS, encoding="utf-8-sig")
    regressions = pd.read_csv(PITCHER_RESULTS, encoding="utf-8-sig")
    sample = pitchers.loc[pitchers["low_sample_flag"].eq(0)].copy()

    summary_rows = []
    for row in regressions.itertuples():
        summary_rows.append(
            {
                "metric": f"Velocity coefficient: {row.outcome} ({row.specification})",
                "value": row.velocity_coef,
                "p_value": row.velocity_p_value,
                "notes": f"player-clustered OLS; n={row.n_obs}",
            }
        )
    summary = pd.DataFrame(
        [
            {"metric": "Pitcher-season observations", "value": len(sample), "p_value": np.nan,
             "notes": "low-sample rows excluded"},
            {"metric": "Unique pitchers", "value": sample["player_id"].nunique(), "p_value": np.nan,
             "notes": "player clusters"},
        ] + summary_rows
    )

    sample["velocity_bucket"] = pd.cut(
        sample["fastball_velo_mph"],
        bins=[-np.inf, 92, 94, 96, np.inf],
        labels=BUCKET_LABELS,
        right=False,
    )
    buckets = (
        sample.groupby("velocity_bucket", observed=False)
        .agg(
            pitcher_seasons=("player_id", "size"),
            average_il_days=("IL_days", "mean"),
            average_war_per_ip=("WAR_per_IP", "mean"),
        )
        .reset_index()
    )

    team_display = team[[
        "season", "total_adjusted_payroll", "CBT_Tax_Paid", "Injured_Salary_Loss",
        "Revenue", "Operating_Income", "Team_Value", "Wins", "Superstar_Era",
        "Pure_Marginal_Benefit",
    ]].copy()
    before = team.loc[team["Superstar_Era"].eq(0)]
    after = team.loc[team["Superstar_Era"].eq(1)]
    comparison_metrics = [
        "Revenue", "Operating_Income", "Team_Value", "total_adjusted_payroll",
        "CBT_Tax_Paid", "Injured_Salary_Loss", "Postseason_Home_Games",
        "Pure_Marginal_Benefit",
    ]
    comparison = pd.DataFrame(
        {
            "metric": comparison_metrics,
            "before_mean": [before[col].mean() for col in comparison_metrics],
            "after_mean": [after[col].mean() for col in comparison_metrics],
        }
    )
    comparison["change"] = comparison["after_mean"] - comparison["before_mean"]
    comparison["percent_change"] = 100 * comparison["change"] / comparison["before_mean"]

    did_display = did.copy()
    did_display["effect_in_report_unit"] = np.where(
        did_display["outcome"].eq("OI_Margin_Pct"),
        did_display["did_coef"] * 100,
        did_display["did_coef"],
    )
    did_display["report_unit"] = np.where(
        did_display["outcome"].eq("OI_Margin_Pct"), "percentage points", "USD"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Key Results", index=False)
        buckets.to_excel(writer, sheet_name="Velocity Buckets", index=False)
        team_display.to_excel(writer, sheet_name="Team Financials", index=False)
        comparison.to_excel(writer, sheet_name="Before After", index=False)
        did_display.to_excel(writer, sheet_name="DiD Results", index=False)

        for worksheet in writer.book.worksheets:
            header_fill = PatternFill("solid", fgColor="1F4E79")
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for index, column in enumerate(worksheet.columns, start=1):
                max_length = max(len(str(cell.value or "")) for cell in column)
                worksheet.column_dimensions[get_column_letter(index)].width = min(max_length + 2, 42)

    print(f"Saved {OUTPUT.relative_to(ROOT)}")
    print("Velocity bucket counts and all means were calculated from data at export time.")
    margin_rows = did_display.loc[did_display["outcome"].eq("OI_Margin_Pct")]
    for row in margin_rows.itertuples():
        print(f"OI margin ({row.sample}): {row.effect_in_report_unit:+.2f} percentage points")


if __name__ == "__main__":
    main()
